"""Local document text extractors. Failures return structured errors, never crash."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Optional

SECRET_LINE_RE = re.compile(
    r"(api[_-]?key|secret|password|token|private[_-]?key|authorization)\s*[:=]",
    re.IGNORECASE,
)

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".html",
    ".htm",
    ".css",
    ".yaml",
    ".yml",
    ".xml",
    ".log",
    ".ini",
    ".cfg",
    ".toml",
    ".rs",
    ".go",
    ".java",
    ".c",
    ".cpp",
    ".h",
    ".cs",
}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff"}


def redact_secret_lines(text: str) -> str:
    lines = []
    for line in text.splitlines():
        if SECRET_LINE_RE.search(line):
            lines.append("[REDACTED]")
        else:
            lines.append(line)
    return "\n".join(lines)


def extract_text(path: Path, max_chars: int = 50000) -> dict:
    """
    Extract text from a local file.
    Returns {ok, text, extractor, error?}.
    """
    ext = path.suffix.lower()
    try:
        if ext in TEXT_EXTENSIONS or ext == "":
            text, extractor = _read_text(path)
        elif ext == ".pdf":
            text, extractor = _read_pdf(path)
        elif ext == ".docx":
            text, extractor = _read_docx(path)
        elif ext in {".xlsx", ".xlsm"}:
            text, extractor = _read_xlsx(path)
        elif ext == ".pptx":
            text, extractor = _read_pptx(path)
        elif ext in IMAGE_EXTENSIONS:
            return {
                "ok": True,
                "text": f"[image file: {path.name}]",
                "extractor": "image_metadata",
            }
        else:
            text, extractor = _read_text(path)
        text = redact_secret_lines(text or "")
        if len(text) > max_chars:
            text = text[:max_chars] + "\n[truncated]"
        return {"ok": True, "text": text, "extractor": extractor}
    except Exception as e:
        return {"ok": False, "text": "", "extractor": ext, "error": str(e)}


def _read_text(path: Path) -> tuple[str, str]:
    data = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding), "text"
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "text"


def _read_pdf(path: Path) -> tuple[str, str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts), "pdf"


def _read_docx(path: Path) -> tuple[str, str]:
    import docx

    doc = docx.Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text), "docx"


def _read_xlsx(path: Path) -> tuple[str, str]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append("\t".join(cells))
    wb.close()
    return "\n".join(parts), "xlsx"


def _read_pptx(path: Path) -> tuple[str, str]:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"# Slide {i}")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
    return "\n".join(parts), "pptx"
